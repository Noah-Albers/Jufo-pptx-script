from pptx.dml.color import RGBColor
from pptx.shapes.placeholder import SlidePlaceholder
from pptx.shapes.autoshape import Shape as AutoShape
from pptx.util import Pt

class EasyTextformatterList:
    def __init__(self, elements):
        self.elements = elements

    def __str__(self):
        return ''.join(self.elements)


class EasyTextformatter:

    def __init__(self, text: str):
        self._text = text
        self._bold = False
        self._italic = False
        self._underline = False
        self._strike = False
        self._subscript = False
        self._superscript = False
        self._kerning = None
        self._shadow = False
        self._emboss = False
        self._engrave = False
        self._font_name = None
        self._font_size = None
        self._font_color = None

    # region Special values

    def kerning(self, kerning: int):
        self._kerning = kerning
        return self

    def font_name(self, font_name: str):
        self._font_name = font_name
        return self

    def font_size(self, size_in_pts: int):
        self._font_size = Pt(size_in_pts)
        return self

    def rgb_color(self, red: int, green: int, blue: int):
        """
        Args:
            red: 0-255
            green: 0-255
            blue: 0-255
        """
        self._font_color = RGBColor(red, green, blue)
        return self

    def hex_color(self, hex: int):
        """
        Args:
            hex: 0x000000 - 0xffffff
            8bit red, 8bit green, 8bit blue
        """
        self._font_color = RGBColor(
            hex >> 16 & 0xff,
            hex >> 8 & 0xff,
            hex & 0xff
        )
        return self

    # endregion

    # region Property's

    @property
    def bold(self):
        self._bold = True
        return self

    @property
    def italic(self):
        self._italic = True
        return self

    @property
    def underline(self):
        self._underline = True
        return self

    @property
    def strike(self):
        self._strike = True
        return self

    @property
    def subscript(self):
        self._subscript = True
        return self

    @property
    def superscript(self):
        self._superscript = True
        return self

    @property
    def shadow(self):
        self._shadow = True
        return self

    @property
    def emboss(self):
        self._emboss = True
        return self

    @property
    def engrave(self):
        self._engrave = True
        return self

    # endregion

    def __str__(self):
        return self._text

    def _apply_to_frame(self, frame: SlidePlaceholder or AutoShape):
        """
        Takes in the frame (Textelement) to apply the styles to
        """

        # Creates the new run (part of the text to format)
        r = frame.paragraphs[0].add_run()
        r.text = self._text

        # Applies all boolean properties if set
        # Apply boolean attributes
        if self._bold:
            r.font.bold = True
        if self._italic:
            r.font.italic = True
        if self._underline:
            r.font.underline = True
        if self._strike:
            r.font.strike = True
        if self._subscript:
            r.font.subscript = True
        if self._superscript:
            r.font.superscript = True
        if self._shadow:
            r.font.shadow = True
        if self._emboss:
            r.font.emboss = True
        if self._engrave:
            r.font.engrave = True

        # Apply non-None attributes
        if self._kerning is not None:
            r.font.kerning = self._kerning
        if self._font_name is not None:
            r.font.name = self._font_name
        if self._font_size is not None:
            r.font.size = self._font_size
        if self._font_color is not None:
            r.font.color.rgb = self._font_color
